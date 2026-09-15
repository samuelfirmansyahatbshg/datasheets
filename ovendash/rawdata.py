"""
The raw data logger file: chart series, and the three values Power Automate
could not compute.

Voltage, Preheat Energy and Total Energy are INDEX/MATCH lookups into a
7,800-row log. Power Automate cannot read an uploaded file's contents, so those
three stayed blank in the flow. Here the file is in hand, so they are ported
directly from `PCGDEOV-22325-6 Whole Chicken.xlsx`:

    Voltage        = INDEX(RawData!C:C, MATCH(3.01, RawData!B:B, 1))
    Preheat Energy = INDEX(RawData!E:E, MATCH(preheat_decimal_min, RawData!B:B, 1))
    Total Energy   = LOOKUP(2, 1/(RawData!E:E<>""), RawData!E:E)

MATCH with type 1 means "last row whose value is <= the target, assuming
ascending order" - not nearest. `_match_le` implements that exactly; treating
it as nearest would silently shift the reading by a row.

The uploaded log is an .xlsx in every sample seen so far despite the form
calling it "Raw Data CSV", so both are handled.
"""

from __future__ import annotations

import csv
import io

# Column positions in the log, 0-indexed: B=time, C=voltage, E=energy.
COL_TIME = 1
COL_VOLTAGE = 2
COL_ENERGY = 4

VOLTAGE_PROBE_MARK = 3.01   # the fixed time marker the workbook reads voltage at

# Enough to draw a smooth trace without shipping ~7,800 points to the browser.
MAX_CHART_POINTS = 900

SERIES_COLORS = ["#c0392b", "#1f5f8b", "#b8860b", "#2d7d5a", "#7b4397", "#c2571a"]


def _to_float(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _rows_from_xlsx(data: bytes) -> list:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _rows_from_csv(data: bytes) -> list:
    text = data.decode("utf-8-sig", errors="replace")
    return [r for r in csv.reader(io.StringIO(text))]


def parse_log(data: bytes, filename: str = "") -> list:
    """Raw bytes -> list of rows, from either format."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        return _rows_from_xlsx(data)
    if filename.lower().endswith(".csv"):
        return _rows_from_csv(data)
    # No usable extension: try xlsx (the observed case), fall back to csv.
    try:
        return _rows_from_xlsx(data)
    except Exception:
        return _rows_from_csv(data)


def _cell(row: list, idx: int):
    return row[idx] if idx < len(row) else None


def _match_le(rows: list, target: float):
    """Excel MATCH(target, col, 1): the last row where time <= target.

    Returns the row itself, or None if every timestamp is already past target.
    """
    found = None
    for row in rows:
        t = _to_float(_cell(row, COL_TIME))
        if t is None:
            continue
        if t <= target:
            found = row
        else:
            break
    return found


def derive_energy(rows: list, preheat_minutes: float | None = None) -> dict:
    """The three values the flow could not compute."""
    out = {"voltage": None, "preheat_energy": None, "total_energy": None}

    row = _match_le(rows, VOLTAGE_PROBE_MARK)
    if row is not None:
        out["voltage"] = _to_float(_cell(row, COL_VOLTAGE))

    if preheat_minutes is not None:
        row = _match_le(rows, preheat_minutes)
        if row is not None:
            out["preheat_energy"] = _to_float(_cell(row, COL_ENERGY))

    # LOOKUP(2, 1/(E:E<>""), E:E) is Excel's idiom for "last non-blank in E".
    for row in reversed(rows):
        v = _to_float(_cell(row, COL_ENERGY))
        if v is not None:
            out["total_energy"] = v
            break

    if out["preheat_energy"] is not None and out["total_energy"] is not None:
        out["cooking_energy"] = out["total_energy"] - out["preheat_energy"]

    return out


def _header_row(rows: list) -> int:
    """Index of the header row.

    Not assumed to be row 0 - these logs sometimes open with a summary block,
    the same trap that once gave a sheet columns named after a float. The
    header is the first row whose time column is *not* numeric but whose
    following row's is.
    """
    for i, row in enumerate(rows[:40]):
        if _to_float(_cell(row, COL_TIME)) is not None:
            return i - 1 if i else 0
    return 0


def _looks_like_counter(points: list) -> bool:
    """True for a plain row-number column.

    These logs lead with an index column, which is numeric and would otherwise
    be charted as a thermocouple - a straight diagonal climbing to several
    thousand, flattening every real trace against the axis. Detected by the
    step between consecutive values being constant and integral.
    """
    if len(points) < 3:
        return False
    steps = {round(b[1] - a[1], 6) for a, b in zip(points, points[1:])}
    if len(steps) != 1:
        return False
    step = steps.pop()
    return step > 0 and abs(step - round(step)) < 1e-9


def build_series(rows: list) -> list:
    """Chart series: every numeric column plotted against time.

    Column B is the X axis. Voltage/energy go on the power axis, everything
    else is treated as a temperature probe - which matches these logs, where
    the remaining numeric columns are thermocouples.

    Counter columns are included but start switched off, so they stay
    available without distorting the default view.
    """
    head_idx = _header_row(rows)
    header = rows[head_idx] if head_idx < len(rows) else []
    data = [r for r in rows[head_idx + 1:] if _to_float(_cell(r, COL_TIME)) is not None]
    if not data:
        return []

    step = max(1, len(data) // MAX_CHART_POINTS)
    sampled = data[::step]

    width = max(len(r) for r in data)
    series = []
    for col in range(width):
        if col == COL_TIME:
            continue

        points = []
        for row in sampled:
            x = _to_float(_cell(row, COL_TIME))
            y = _to_float(_cell(row, col))
            if x is not None and y is not None:
                points.append([round(x, 3), round(y, 2)])
        if len(points) < 2:
            continue

        name = ""
        if col < len(header) and header[col] not in (None, ""):
            name = str(header[col]).strip()
        if not name:
            name = f"Column {col + 1}"

        axis = "power" if col in (COL_VOLTAGE, COL_ENERGY) else "temp"
        series.append(
            {
                "name": name,
                "axis": axis,
                "color": SERIES_COLORS[len(series) % len(SERIES_COLORS)],
                "points": points,
                "default_on": not _looks_like_counter(points),
            }
        )

    return series


def attach_raw_data(submission, source) -> None:
    """Download the log and hang the chart series + energy values off the submission.

    A no-op when the file cannot be fetched, so the dashboard still renders
    everything else - the chart panel explains its own absence.
    """
    att = submission.attachments.get("RawDataCSV")
    downloader = getattr(source, "download", None)
    if not att or not att.present or downloader is None:
        return

    data = downloader(att.drive_id, att.item_id)
    rows = parse_log(data, att.name)
    if not rows:
        return

    submission.raw_series = build_series(rows)

    preheat_minutes = None
    raw = submission.info.get("PreheatTime") or submission.info.get("PreheatTimeDecimal")
    if raw:
        preheat_minutes = _to_float(raw)
    submission.energy = derive_energy(rows, preheat_minutes)
